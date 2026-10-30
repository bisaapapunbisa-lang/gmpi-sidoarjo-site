from docx import Document
from docx.shared import Cm, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree
import copy

RED=RGBColor(0xE5,0x1C,0x24); DRED=RGBColor(0x9A,0x29,0x28); INK=RGBColor(0x15,0x15,0x13); GREY=RGBColor(0x5F,0x5B,0x53); WHITE=RGBColor(255,255,255); GOLD=RGBColor(0xA8,0x80,0x2A)
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.left_margin=sec.right_margin=Cm(2.2)
sec.top_margin=Cm(5.2); sec.bottom_margin=Cm(3.2)
sec.header_distance=Cm(0.9); sec.footer_distance=Cm(0.55)

st=doc.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
st.paragraph_format.space_after=Pt(0); st.paragraph_format.line_spacing=1.15

def run(p,text,size=None,bold=False,color=None,italic=False,font=None,spacing=None,caps=False):
    r=p.add_run(text); f=r.font
    if size: f.size=Pt(size)
    f.bold=bold; f.italic=italic
    if color is not None: f.color.rgb=color
    if font:
        f.name=font; r._element.rPr.rFonts.set(qn('w:eastAsia'),font); r._element.rPr.rFonts.set(qn('w:cs'),font)
    if spacing is not None:
        sp=OxmlElement('w:spacing'); sp.set(qn('w:val'),str(int(spacing*20))); r._element.rPr.append(sp)
    if caps: f.all_caps=True
    return r

_id=[100]
def anchor(run_, path, width_cm, x_cm, y_cm, rel_h='page', rel_v='page', behind=True, height_cm=None):
    """Sisipkan gambar mengambang (di belakang teks) pada posisi absolut halaman."""
    if height_cm: run_.add_picture(path, width=Cm(width_cm), height=Cm(height_cm))
    else: run_.add_picture(path, width=Cm(width_cm))
    drawing=run_._element.findall(qn('w:drawing'))[-1]
    inline=drawing.find(qn('wp:inline'))
    _id[0]+=1
    anc=OxmlElement('wp:anchor')
    for k,v in dict(distT='0',distB='0',distL='0',distR='0',simplePos='0',relativeHeight=str(_id[0]),behindDoc='1' if behind else '0',locked='1',layoutInCell='1',allowOverlap='1').items(): anc.set(k,v)
    sp=OxmlElement('wp:simplePos'); sp.set('x','0'); sp.set('y','0'); anc.append(sp)
    for tag,rel,val in (('wp:positionH',rel_h,x_cm),('wp:positionV',rel_v,y_cm)):
        e=OxmlElement(tag); e.set('relativeFrom',rel); o=OxmlElement('wp:posOffset'); o.text=str(int(Cm(val))); e.append(o); anc.append(e)
    anc.append(copy.deepcopy(inline.find(qn('wp:extent'))))
    ee=OxmlElement('wp:effectExtent')
    for k in 'lrtb': ee.set(k,'0')
    anc.append(ee); anc.append(OxmlElement('wp:wrapNone'))
    dp=copy.deepcopy(inline.find(qn('wp:docPr'))); dp.set('id',str(_id[0])); anc.append(dp)
    anc.append(copy.deepcopy(inline.find(qn('wp:cNvGraphicFramePr'))))
    anc.append(copy.deepcopy(inline.find(qn('a:graphic'))))
    drawing.remove(inline); drawing.append(anc)


def fix_widths(tbl, widths_cm):
    t=tbl._tbl; tblPr=t.tblPr
    lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    tw=tblPr.find(qn('w:tblW'))
    if tw is None: tw=OxmlElement('w:tblW'); tblPr.append(tw)
    tw.set(qn('w:w'),str(int(sum(widths_cm)*567))); tw.set(qn('w:type'),'dxa')
    ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),'0'); ind.set(qn('w:type'),'dxa'); tblPr.append(ind)
    cm=OxmlElement('w:tblCellMar')
    for side in ('left','right'):
        e=OxmlElement(f'w:{side}'); e.set(qn('w:w'),'0'); e.set(qn('w:type'),'dxa'); cm.append(e)
    tblPr.append(cm)
    grid=t.tblGrid
    for gc,w in zip(grid.findall(qn('w:gridCol')),widths_cm): gc.set(qn('w:w'),str(int(w*567)))
    for row in tbl.rows:
        for c,w in zip(row.cells,widths_cm): c.width=Cm(w)

def no_borders(tbl):
    tblPr=tbl._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement(f'w:{side}'); e.set(qn('w:val'),'nil'); b.append(e)
    tblPr.append(b)
def cell_margins(cell, **kw):
    tcPr=cell._tc.get_or_add_tcPr(); m=OxmlElement('w:tcMar')
    for side,v in kw.items():
        e=OxmlElement(f'w:{side}'); e.set(qn('w:w'),str(int(v*567))); e.set(qn('w:type'),'dxa'); m.append(e)
    tcPr.append(m)
def vcenter(cell):
    v=OxmlElement('w:vAlign'); v.set(qn('w:val'),'center'); cell._tc.get_or_add_tcPr().append(v)

# ================= KOP (header) =================
hdr=sec.header; hdr.is_linked_to_previous=False
hp=hdr.paragraphs[0]
r0=hp.add_run()
anchor(r0,'ornamen-atas.png',21,0,0)                 # ornamen kanan atas
anchor(r0,'watermark.png',12.4,4.3,9.2)             # logo pudar di tengah halaman
anchor(r0,'pita-bawah.png',21,0,29.7-2.6)            # pita bawah
anchor(r0,'logo.png',2.8,2.15,1.2,behind=False,height_cm=3.018)   # logo
anchor(r0,'garis-kop.png',16.6,2.2,4.42)            # garis ganda bawah kop
hp.paragraph_format.space_after=Pt(0)
fmt=hp.paragraph_format; hp.runs[0].font.size=Pt(1)

tbl=hdr.add_table(rows=1, cols=2, width=Cm(16.6))
tbl.alignment=WD_TABLE_ALIGNMENT.LEFT; no_borders(tbl); tbl.autofit=False
c0,c1=tbl.rows[0].cells
fix_widths(tbl,[3.3,13.3])
tbl.rows[0].height=Cm(2.9)
for c in (c0,c1): vcenter(c)
cell_margins(c0,left=0,right=0.15,top=0,bottom=0); cell_margins(c1,left=0.2,right=0,top=0,bottom=0)
p=c0.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT
p=c1.paragraphs[0]
run(p,'GERAKAN MILITANSI PEJUANG INDONESIA',10.5,True,INK,spacing=0.6)
p=c1.add_paragraph(); p.paragraph_format.space_before=Pt(1)
run(p,'DEWAN PIMPINAN DAERAH',8,True,DRED,spacing=2.2)
p=c1.add_paragraph(); p.paragraph_format.line_spacing=0.95
run(p,'GMPI ',26,True,RED,spacing=0.5)
run(p,'KABUPATEN SIDOARJO',18,True,INK,spacing=0.3)
p=c1.add_paragraph(); p.paragraph_format.space_before=Pt(3)
run(p,'Badan Hukum: SK Menkumham RI No. AHU-0010942.AH.01.07.TAHUN 2021',7.5,False,GREY)
p=c1.add_paragraph()
run(p,'NPWP 53.380.197.3-408.000',7.5,False,GREY)
run(p,'   ★   ',7.5,True,GOLD)
run(p,'Setia Bersaudara Sampai Urat Nadi',7.5,True,DRED,italic=True)

# ================= Pita bawah (footer) =================
ftr=sec.footer; ftr.is_linked_to_previous=False
fp=ftr.paragraphs[0]; fp.paragraph_format.left_indent=Cm(-0.6)
run(fp,'SEKRETARIAT DPD GMPI KABUPATEN SIDOARJO',8,True,WHITE,spacing=1.2)
fp2=ftr.add_paragraph(); fp2.paragraph_format.left_indent=Cm(-0.6); fp2.paragraph_format.space_before=Pt(2)
run(fp2,'Alamat: ',7.5,True,WHITE); run(fp2,'Perum, Jl. Graha Bougenville Kav. Depag No.31, Kemiri, Kec. Sidoarjo, Kab. Sidoarjo, Jawa Timur 61234',7.5,False,WHITE)
apJ=alamat_join=None
fp3=ftr.add_paragraph(); fp3.paragraph_format.left_indent=Cm(-0.6); fp3.paragraph_format.space_before=Pt(1)
run(fp3,'Telp/WA: ',7.5,True,WHITE); run(fp3,'+62 812-3591-1689',7.5,False,WHITE)
run(fp3,'    •    ',7.5,True,WHITE)
run(fp3,'Email: ',7.5,True,WHITE); run(fp3,'gmpinusantara@gmail.com',7.5,False,WHITE)


def para(text='',size=11,bold=False,align=None,before=0,after=0,indent=None,italic=False,color=None):
    p=doc.add_paragraph()
    if text: run(p,text,size,bold,color,italic)
    if align: p.alignment=align
    p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
    if indent is not None: p.paragraph_format.first_line_indent=Cm(indent)
    return p

t=doc.add_table(rows=3,cols=2); no_borders(t); t.autofit=False
rows=[('Nomor',': …… /DPD-GMPI/SDA/…… /2026'),('Lampiran',': -'),('Perihal',': Undangan Rapat Koordinasi & Penguatan Instruksi Teknis')]
for i,(a,b) in enumerate(rows):
    c=t.rows[i].cells
    run(c[0].paragraphs[0],a,11); run(c[1].paragraphs[0],b,11,bold=(a=='Perihal'))
    for cc in c: cell_margins(cc,left=0,right=0,top=0,bottom=0)
fix_widths(t,[2.6,11])
dp=para('Sidoarjo, …… Oktober 2026',align=WD_ALIGN_PARAGRAPH.RIGHT,before=6)
t._tbl.addprevious(dp._element)

para('Kepada Yth.',before=12)
para('Para Pimpinan PAC & DPC',bold=True)
para('GMPI se-Kabupaten Sidoarjo',bold=True)
para('di –')
p=para('Tempat'); p.paragraph_format.left_indent=Cm(1)

para('Dengan hormat,',before=12,after=6)
para('Salam Militansi! Semoga Bapak/Ibu/Saudara senantiasa dalam lindungan Tuhan Yang Maha Esa.',align=WD_ALIGN_PARAGRAPH.JUSTIFY,after=6,indent=1)
para('Dalam rangka penguatan barisan dan konsolidasi kegiatan Dewan Pimpinan Daerah, kami mengundang Bapak/Ibu/Saudara hadir dalam Rapat Koordinasi & Penguatan Instruksi Teknis yang Insya Allah akan diselenggarakan pada:',align=WD_ALIGN_PARAGRAPH.JUSTIFY,after=8,indent=1)

a=doc.add_table(rows=4,cols=2); no_borders(a); a.autofit=False
jadwal=[('Hari',':  Rabu'),('Tanggal',':  21 Oktober 2026'),('Pukul',':  18.30 WIB – selesai'),('Tempat',':  Graha Bougenville Kav. 22–23, Kemiri, Kec. Sidoarjo')]
for i,(k_,v) in enumerate(jadwal):
    c=a.rows[i].cells
    rp=run(c[0].paragraphs[0],k_,11,True)
    run(c[1].paragraphs[0],v,11)
    for cc in c: cell_margins(cc,left=0,right=0,top=0.02,bottom=0.02)
fix_widths(a,[2.6,12])
for p in a.rows[0].cells: pass
# geser indent tabel jadwal sedikit ke kanan
tblPr=a._tbl.tblPr; ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),str(int(1*567))); ind.set(qn('w:type'),'dxa'); tblPr.append(ind)

para('Mengingat pentingnya rapat ini, kehadiran Bapak/Ibu/Saudara kami harapkan tepat waktu. Konfirmasi kehadiran mohon diisi melalui undangan digital kami: https://www.gmpisidoarjo.com/rsvp',align=WD_ALIGN_PARAGRAPH.JUSTIFY,before=8,after=6,indent=1)
para('Demikian undangan ini kami sampaikan. Atas perhatian dan kehadiran Bapak/Ibu/Saudara, kami ucapkan terima kasih.',align=WD_ALIGN_PARAGRAPH.JUSTIFY,after=6,indent=1)

para('Setia Bersaudara Sampai Urat Nadi!',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,before=8)
para('DEWAN PIMPINAN DAERAH',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,before=14,size=10.5)
para('GERAKAN MILITANSI PEJUANG INDONESIA KABUPATEN SIDOARJO',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,size=10.5)
s=doc.add_table(rows=1,cols=2); no_borders(s); s.autofit=False
fix_widths(s,[8.3,8.3])
ck,cs_=s.rows[0].cells
p=ck.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(6)
run(p,'Ketua',11,True)
pt=ck.add_paragraph(); pt.alignment=WD_ALIGN_PARAGRAPH.CENTER; pt.paragraph_format.space_before=Pt(0); pt.paragraph_format.space_after=Pt(0)
pt.add_run().add_picture('ttd-ketua.png',width=Cm(2.7))
qn_=ck.add_paragraph(); qn_.alignment=WD_ALIGN_PARAGRAPH.CENTER; qn_.paragraph_format.space_before=Pt(0)
run(qn_,'( HASMIN MARJAN )',11,True)
p=cs_.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(6)
run(p,'Sekretaris',11,True)
for _ in range(3): q=cs_.add_paragraph(); q.paragraph_format.space_after=Pt(8)
q=cs_.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
run(q,'( ……………………………… )',11,True)

doc.core_properties.title='Undangan Rapat DPD GMPI Kab. Sidoarjo'
doc.core_properties.author='DPD GMPI Kabupaten Sidoarjo'
doc.save('Surat-Undangan-Rapat-21Okt2026.docx')
print('saved')
